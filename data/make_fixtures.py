"""Generate DOCX and PDF versions of league_rules.md for testing uploads."""
from pathlib import Path

from docx import Document as Docx
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

HERE = Path(__file__).parent
lines = (HERE / "league_rules.md").read_text(encoding="utf-8").splitlines()

# ---- DOCX ----
d = Docx()
for ln in lines:
    if ln.startswith("# "):
        d.add_heading(ln[2:], level=0)
    elif ln.startswith("## "):
        d.add_heading(ln[3:], level=1)
    elif ln.strip():
        d.add_paragraph(ln.lstrip("- ").strip())
d.save(HERE / "league_rules.docx")

# ---- PDF ----
styles = getSampleStyleSheet()
story = []
for ln in lines:
    if ln.startswith("# "):
        story.append(Paragraph(ln[2:], styles["Title"]))
    elif ln.startswith("## "):
        story.append(Paragraph(ln[3:], styles["Heading2"]))
    elif ln.strip():
        story.append(Paragraph(ln.lstrip("- ").strip(), styles["Normal"]))
    story.append(Spacer(1, 4))
SimpleDocTemplate(str(HERE / "league_rules.pdf"), pagesize=letter).build(story)

print("Wrote league_rules.docx and league_rules.pdf")
